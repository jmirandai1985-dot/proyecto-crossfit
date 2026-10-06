"""Ventas del Bazar: UNA definición para el BI, el Excel de Reportes y el historial del alumno.

Diagnóstico (2026-10)
---------------------
La misma pregunta ("¿cuánto vendió el Bazar?") tenía TRES respuestas distintas:

  · **BI** (`daily_kpis.ingresos_bazar`, la tarjeta "Ingresos por bazar" del KPI diario):
    sumaba `transacciones_financieras` con `categoria='bazar'`. Como el Bazar NO inserta
    transacciones, el número era SIEMPRE 0: la venta del Bazar era invisible para el BI.
  · **Excel de Reportes** ("Bazar / Tienda" y la pestaña "Bazar y Servicios"): contaba
    `estado != 'cancelado'`, o sea que un pedido con el comprobante SIN revisar (`pendiente`)
    ya se publicaba como venta.
  · **Historial del alumno** (pestaña Pagos): `validado`/`entregado` — la única de las tres
    que estaba bien.

Y `monthly_kpis` no tenía ninguna columna de Bazar: la pestaña Mensual publicaba "Ingresos del
mes" sin el Bazar y sin forma de verlo.

Lo que fija este archivo
------------------------
  A. PURA (sin BD): la lista de estados cobrados es UNA
     (`shared.estados.ESTADOS_PAGO_BAZAR`, la misma que re-exporta `app.core.estados`).
  B. CASO ARMADO contra TEST (escribe y RESTAURA): un alumno temporal con 5 pedidos de un mes
     sin datos reales — validado, entregado, pendiente, cancelado y uno cargado a las 22:30 de
     Chile del último día del mes (01:30 UTC del día siguiente) — y las comprobaciones de que
     los consumidores dicen LO MISMO con esa data:
       · `metricas_service.ventas_bazar` (la definición única),
       · el BI (día y mes, con `daily_kpis`/`monthly_kpis`/`GET /kpis/mensual`),
       · el Excel de Reportes (detalle, total y archivo),
       · la pestaña Bazar del historial del alumno (datos del retiro).
  C. VIVO (sin escribir): `/reportes/` publica `ventasBazar` con la MISMA métrica.

No usa el servidor: habla con la MISMA rama TEST por `SessionLocal` (igual que
`test_historial_alumno_service.py`). El guard `is_test_db_url` falla CERRADO: si el proceso no
está apuntando a TEST, no corre nada.

Correr:
    ENVIRONMENT=test py -3.12 -m pytest tests/test_bazar_ventas_compartidas.py -q
    (o dentro del contenedor: docker compose exec -T backend python -m pytest \
        tests/test_bazar_ventas_compartidas.py -q)
"""
import calendar
import secrets
import sys
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from sqlalchemy import text

_BACKEND = Path(__file__).resolve().parents[1]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from app.services import metricas_service as metricas          # noqa: E402
from app.utils.santiago import hoy_santiago                     # noqa: E402

TENANT_ID = 1

# ── El mes del caso armado: 2020-01 no tiene NINGÚN dato real, así que los números de abajo
#    son sólo de este test (y se borran al terminar).
ANIO, MES = 2020, 1
VALIDADO = 20000
ENTREGADO = 5000
MADRUGA = 777          # validado, cargado 01:30 UTC del 01/02 = 22:30 del 31/01 en Chile
PENDIENTE = 9999
CANCELADO = 500
COBRADO_MES = VALIDADO + ENTREGADO + MADRUGA


@pytest.fixture(scope="module")
def db():
    """Sesión contra la MISMA rama TEST que ve la API (falla cerrado)."""
    from app.core.config import is_test_db_url, settings
    from app.db.database import SessionLocal

    if not is_test_db_url(settings.DATABASE_URL):
        pytest.fail("DATABASE_URL no es una rama de TEST: aborto por seguridad "
                    "(define ENVIRONMENT=test / revisa .env.test)")
    print("\n[OK] base de datos de TEST confirmada\n")
    session = SessionLocal()
    yield session
    session.close()


# ══════════════════════════════════════════════════════════════════════════════
# A. Reglas puras (sin BD)
# ══════════════════════════════════════════════════════════════════════════════
def test_a1_la_lista_de_pago_del_bazar_es_una_sola_para_todos():
    """`validado`/`entregado` viven UNA vez y los dos lados leen ESA constante."""
    from app.core import estados as app_estados
    from shared import estados as shared_estados

    assert shared_estados.ESTADOS_PAGO_BAZAR == ("validado", "entregado")
    assert shared_estados.lista_sql_pago_bazar() == "'validado', 'entregado'"
    # La app re-exporta el MISMO objeto: no hay una segunda lista que pueda quedar vieja.
    assert app_estados.ESTADOS_PAGO_BAZAR is shared_estados.ESTADOS_PAGO_BAZAR


def test_a2_el_predicado_orm_arma_el_in_con_la_lista_compartida():
    """`pago_bazar(Pedido.estado)` compila a `estado IN ('validado', 'entregado')`."""
    from app.core.estados import pago_bazar
    from app.models.pedido import Pedido

    sql = str(pago_bazar(Pedido.estado).compile(compile_kwargs={"literal_binds": True}))

    assert "IN ('validado', 'entregado')" in sql


# ══════════════════════════════════════════════════════════════════════════════
# B. Caso armado contra TEST (escribe y RESTAURA)
# ══════════════════════════════════════════════════════════════════════════════
# Alfabeto del código de retiro (el MISMO del CHECK de la migración 044: sin 0/O ni 1/I/L).
ALFABETO_CODIGO = "23456789ABCDEFGHJKMNPQRSTUVWXYZ"


def _utc(dia, hora, minuto) -> datetime:
    """Instante UTC del mes del caso (los `timestamptz` de la BD se guardan en UTC)."""
    return datetime(ANIO, MES, dia, hora, minuto, tzinfo=timezone.utc)


def _codigo_libre(db) -> str:
    """Un `UB-XXXX` que el box de TEST no tenga (el índice único es por tenant)."""
    for _ in range(20):
        cand = "UB-" + "".join(secrets.choice(ALFABETO_CODIGO) for _ in range(4))
        ocupado = db.execute(text(
            "SELECT 1 FROM pedidos WHERE tenant_id = :t AND codigo_retiro = :c"),
            {"t": TENANT_ID, "c": cand}).first()
        if not ocupado:
            return cand
    pytest.fail("no se pudo generar un código de retiro libre en TEST")


@pytest.fixture
def caso(db):
    """Alumno temporal con 5 pedidos del mes de prueba; borra TODO lo que crea al salir.

    `madruga` está cargado como `2020-02-01 01:30+00`: en Chile es el 31/01 a las 22:30, así que
    cuenta en ENERO. Es el caso que atrapa al SQL que compara `fecha_pedido` contra los bordes del
    día en UTC (esa venta se le iba al mes siguiente).
    """
    sufijo = f"{datetime.now():%Y%m%d%H%M%S}"
    # La sesión es la MISMA para todo el módulo: si un test anterior la dejó en una transacción
    # abortada (Postgres la marca así hasta el rollback), acá se limpia para no arrastrar el error.
    db.rollback()
    producto = db.execute(text(
        "SELECT id, nombre FROM productos WHERE tenant_id = :t ORDER BY id LIMIT 1"),
        {"t": TENANT_ID}).first()
    if not producto:
        pytest.skip("TEST no tiene productos para el caso armado")
    # `rol` es el enum `rol_usuario` (alumno | coach | administrador): "admin" NO existe como
    # etiqueta y Postgres rechaza el IN entero.
    admin = db.execute(text(
        "SELECT id, nombre FROM usuarios WHERE tenant_id = :t "
        "AND rol = 'administrador' ORDER BY id LIMIT 1"),
        {"t": TENANT_ID}).first()
    if not admin:
        pytest.skip("TEST no tiene un admin para el caso armado")

    codigo = _codigo_libre(db)
    alumno_id = db.execute(text("""
        INSERT INTO usuarios (tenant_id, rut, nombre, correo, password_hash, rol, activo,
                              estado, created_at)
        VALUES (:t, :r, 'Alumno Bazar TEST', :c, 'x', 'alumno', true, 'activo', :alta)
        RETURNING id
    """), {"t": TENANT_ID, "r": f"97{sufijo[-8:]}-7", "c": f"bazar.{sufijo}@test.local",
          "alta": datetime(ANIO, MES, 2, 12, 0, tzinfo=timezone.utc)}).scalar()

    def _pedido(estado, total, fecha_utc, codigo_retiro=None, entregado_en=None):
        return db.execute(text("""
            INSERT INTO pedidos (tenant_id, alumno_id, producto_id, cantidad, total, estado,
                                 fecha_pedido, codigo_retiro, entregado_por, entregado_en)
            VALUES (:t, :a, :p, 1, :m, :e, :f, :c, :ep, :ee) RETURNING id
        """), {"t": TENANT_ID, "a": alumno_id, "p": producto.id, "m": total, "e": estado,
              "f": fecha_utc, "c": codigo_retiro,
              "ep": admin.id if entregado_en else None,
              "ee": entregado_en}).scalar()

    ids = [
        _pedido("validado", VALIDADO, _utc(10, 15, 0), codigo_retiro=codigo),
        _pedido("entregado", ENTREGADO, _utc(20, 15, 0), entregado_en=_utc(21, 14, 0)),
        _pedido("pendiente", PENDIENTE, _utc(15, 15, 0)),
        _pedido("cancelado", CANCELADO, _utc(16, 15, 0)),
        # 01:30 UTC del 01/02 = 22:30 del 31/01 en Chile.
        _pedido("validado", MADRUGA, datetime(ANIO, MES + 1, 1, 1, 30, tzinfo=timezone.utc)),
    ]
    db.commit()

    try:
        yield {"alumno_id": alumno_id, "pedidos": ids, "producto": producto.nombre,
               "codigo": codigo, "admin": admin.nombre,
               "entregado_en": _utc(21, 14, 0)}
    finally:
        db.execute(text("DELETE FROM pedidos WHERE alumno_id = :a"), {"a": alumno_id})
        db.execute(text("DELETE FROM usuarios WHERE id = :a"), {"a": alumno_id})
        # Fotos del BI que creó el caso (el mes 2020-01 no lo usa nadie más).
        db.execute(text("DELETE FROM daily_kpis WHERE tenant_id = :t "
                        "AND fecha >= :i AND fecha <= :f"),
                   {"t": TENANT_ID, "i": date(ANIO, MES, 1), "f": date(ANIO, MES + 1, 1)})
        db.execute(text("DELETE FROM monthly_kpis WHERE tenant_id = :t "
                        "AND year = :y AND month = :m"),
                   {"t": TENANT_ID, "y": ANIO, "m": MES})
        db.commit()


def _fin_de_mes(anio: int, mes: int) -> date:
    return date(anio, mes, calendar.monthrange(anio, mes)[1])


def test_b1_la_metrica_cuenta_lo_cobrado_y_no_lo_pendiente(db, caso):
    """`ventas_bazar` = validado + entregado; el pendiente y el cancelado NO son plata."""
    enero = metricas.ventas_bazar(db, TENANT_ID, date(ANIO, MES, 1), _fin_de_mes(ANIO, MES))

    assert enero == COBRADO_MES
    # El pendiente (que el Excel contaba como venta) y el cancelado quedan afuera.
    assert enero != COBRADO_MES + PENDIENTE
    assert enero != COBRADO_MES + CANCELADO


def test_b2_la_venta_de_la_noche_pertenece_al_dia_chileno(db, caso):
    """22:30 del 31/01 en Chile es ENERO, aunque en UTC ya sea el 01/02."""
    enero = metricas.ventas_bazar(db, TENANT_ID, date(ANIO, MES, 1), _fin_de_mes(ANIO, MES))
    febrero = metricas.ventas_bazar(db, TENANT_ID, date(ANIO, 2, 1), _fin_de_mes(ANIO, 2))

    assert enero == COBRADO_MES, "la venta de las 22:30 del 31/01 tiene que caer en enero"
    assert febrero == 0, "febrero no se puede llevar la venta de la noche del 31/01"


def test_b3_el_bi_diario_deja_de_publicar_bazar_en_cero(db, caso):
    """La tarjeta "Ingresos por bazar" del día sale de los pedidos cobrados (antes: SIEMPRE 0)."""
    from app.api.v1 import kpis_populate
    from app.models.daily_kpis import DailyKpi

    resp = kpis_populate.populate_daily_kpis(fecha=date(ANIO, MES, 31), db=db, _auth=True)

    # El 31/01 (Chile) sólo tiene el pedido de las 22:30.
    assert resp["valores"]["ingresos_bazar"] == MADRUGA
    assert resp["valores"]["ingresos_total"] == (
        resp["valores"]["ingresos_membresia"] + MADRUGA)

    fila = db.query(DailyKpi).filter(DailyKpi.tenant_id == TENANT_ID,
                                     DailyKpi.fecha == date(ANIO, MES, 31)).first()
    assert float(fila.ingresos_bazar) == MADRUGA
    assert float(fila.ingresos_total) == (
        float(fila.ingresos_membresia) + float(fila.ingresos_bazar))


def test_b4_el_bi_mensual_y_la_pestana_mensual_publican_el_bazar(db, caso):
    """`monthly_kpis.ingresos_bazar` + `GET /kpis/mensual` (sin tocar el MRR ni el ingreso)."""
    from app.api.v1 import kpis, kpis_populate
    from app.models.monthly_kpis import MonthlyKpi

    pop = kpis_populate._upsert_mes_monthly(db, TENANT_ID, ANIO, MES)
    assert pop["valores"]["ingresos_bazar"] == COBRADO_MES

    fila = db.query(MonthlyKpi).filter(MonthlyKpi.tenant_id == TENANT_ID,
                                       MonthlyKpi.year == ANIO,
                                       MonthlyKpi.month == MES).first()
    assert float(fila.ingresos_bazar) == COBRADO_MES
    # Agregar la columna NO movió lo que ya estaba: MRR e ingreso del mes son los mismos.
    assert float(fila.mrr) == pop["valores"]["mrr"]
    assert float(fila.ingresos_total) == pop["valores"]["ingresos_total"]

    admin = {"usuario_id": 1, "tenant_id": TENANT_ID, "rol": "administrador"}
    vista = kpis.get_kpis_mensual(year=ANIO, month=MES, db=db, current_user=admin)
    assert vista["ingresos_bazar"] == COBRADO_MES
    # El bazar vive APARTE del ingreso del mes: ese sigue siendo el neto de las transacciones.
    assert vista["ingresos_total"] == metricas.ingresos_netos(
        db, TENANT_ID, date(ANIO, MES, 1), _fin_de_mes(ANIO, MES))

    # El selector de la pestaña Mensual también lo trae.
    per = kpis.get_kpis_mensual_periodos(db=db, current_user=admin)
    fila_per = next(p for p in per["periodos"]
                    if (p["year"], p["month"]) == (ANIO, MES))
    assert fila_per["ingresos_bazar"] == COBRADO_MES


def test_b5_el_excel_lista_y_suma_los_mismos_pedidos_que_la_metrica(db, caso):
    """El detalle del Excel y su TOTAL son los pedidos COBRADOS (el pendiente ya no suma)."""
    from app.services import reportes_service

    inicio, fin = date(ANIO, MES, 1), date(ANIO, MES + 1, 1)   # fin EXCLUSIVO, como el Excel
    filas = reportes_service._pedidos_bazar_mes(db, TENANT_ID, inicio, fin)

    # Antes el detalle listaba también el pendiente (y su fila TOTAL no sumaba con las filas).
    assert {f.estado for f in filas} == {"validado", "entregado"}
    assert sum(float(f.total) for f in filas) == COBRADO_MES
    assert sum(float(f.total) for f in filas) == metricas.ventas_bazar(
        db, TENANT_ID, inicio, _fin_de_mes(ANIO, MES))
    # Los pendientes se cuentan aparte: la tarjeta "Pendientes" del Excel NO es venta.
    assert reportes_service._pedidos_bazar_pendientes_mes(db, TENANT_ID, inicio, fin) == 1


def test_b6_el_archivo_del_excel_se_genera_con_la_pestana_del_bazar(db):
    """El Excel completo (del mes en curso) se arma y trae la pestaña "Bazar y Servicios"."""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services import reportes_service

    hoy = hoy_santiago()
    datos = reportes_service.crear_reporte_ventas_mensual_bytes(
        db, TENANT_ID, hoy.month, hoy.year)

    assert isinstance(datos, bytes) and len(datos) > 5000
    wb = load_workbook(BytesIO(datos))
    assert "Bazar y Servicios" in wb.sheetnames


def test_b7_la_pestana_bazar_del_historial_muestra_el_retiro(db, caso):
    """La pestaña Bazar del alumno: todos los pedidos, con su código y quién/cuándo entregó."""
    from app.services import historial_alumno_service as svc

    datos = svc.panel(db, caso["alumno_id"], TENANT_ID, seccion="bazar")["datos"]
    t = datos["totales"]

    assert t["pedidos"] == 5
    assert t["cobrados"] == 3          # validado + entregado + el de las 22:30
    assert t["entregados"] == 1
    assert t["pendientes"] == 1
    assert t["cobrado_clp"] == COBRADO_MES
    assert t["unidades"] == 5

    validado = next(i for i in datos["items"] if i["id"] == caso["pedidos"][0])
    assert validado["producto"] == caso["producto"]
    assert validado["cantidad"] == 1
    assert validado["total_clp"] == VALIDADO
    assert validado["estado"] == "validado"
    assert validado["cobrado"] is True
    assert validado["codigo_retiro"] == caso["codigo"]
    # Todavía no se entregó: sin usuario ni fecha (no se inventan).
    assert validado["entregado_por"] is None and validado["entregado_en"] is None

    entregado = next(i for i in datos["items"] if i["id"] == caso["pedidos"][1])
    assert entregado["entregado_por"] == caso["admin"]
    assert entregado["entregado_en"] == caso["entregado_en"].date()

    # El pendiente se LISTA (es el historial de pedidos) pero no suma plata ni tiene código.
    pendiente = next(i for i in datos["items"] if i["id"] == caso["pedidos"][2])
    assert pendiente["cobrado"] is False
    assert pendiente["codigo_retiro"] is None


# ══════════════════════════════════════════════════════════════════════════════
# C. Vivo (sin escribir): /reportes/ publica el bazar con la MISMA métrica
# ══════════════════════════════════════════════════════════════════════════════
def test_c1_reportes_vivos_publican_el_bazar_con_la_misma_regla(db):
    """`ventasBazar` es un campo propio de /reportes/ y excluye los pendientes."""
    from app.api.v1 import reportes

    hoy = hoy_santiago()
    admin = {"usuario_id": 1, "tenant_id": TENANT_ID, "rol": "administrador"}
    data = reportes.obtener_reportes_analytics(
        tenant_id=TENANT_ID, db=db, current_user=admin)

    assert "ventasBazar" in data
    inicio, fin = hoy.replace(day=1), _fin_de_mes(hoy.year, hoy.month)
    assert data["ventasBazar"] == metricas.ventas_bazar(db, TENANT_ID, inicio, fin)

    # Comprobación contra SQL escrito aparte: si el endpoint sumara los pendientes (lo que hacía
    # el Excel viejo con `estado != 'cancelado'`), estos dos números no darían igual.
    fila = db.execute(text("""
        SELECT COALESCE(SUM(total) FILTER (WHERE estado IN ('validado', 'entregado')), 0),
               COALESCE(SUM(total) FILTER (WHERE estado = 'pendiente'), 0)
        FROM pedidos
        WHERE tenant_id = :t
          AND (fecha_pedido AT TIME ZONE 'America/Santiago')::date >= :i
          AND (fecha_pedido AT TIME ZONE 'America/Santiago')::date <= :f
    """), {"t": TENANT_ID, "i": inicio, "f": fin}).first()

    assert data["ventasBazar"] == float(fila[0])
    if float(fila[1]) > 0:
        assert data["ventasBazar"] != float(fila[0]) + float(fila[1]), (
            "los pedidos pendientes del mes NO son venta")
